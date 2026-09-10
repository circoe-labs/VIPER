"""Spreadsheet adapters: XLSX/CSV reading, encodings, delimiters, merged cells, limits, errors."""

import io
import zipfile
from datetime import datetime

import pytest
from openpyxl import Workbook as OpenpyxlWorkbook

from app.core.config import Settings
from app.services.imports.diagnostics import DiagnosticCode, ImportRejectedError
from app.services.imports.workbook import (
    OLE_MAGIC,
    FileFormat,
    ImportLimits,
    read_workbook,
)

LIMITS = ImportLimits()
CSV_HEADER = "Entreprise;Nom;Prénom;Mail"


def xlsx(rows: list[list[object]], *, merges: tuple[str, ...] = ()) -> bytes:
    book = OpenpyxlWorkbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Prospects"
    for row in rows:
        sheet.append(row)
    for cell_range in merges:
        sheet.merge_cells(cell_range)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def rejected_code(filename: str, content: bytes, limits: ImportLimits = LIMITS) -> DiagnosticCode:
    with pytest.raises(ImportRejectedError) as caught:
        read_workbook(filename, content, limits)
    return caught.value.code


def test_xlsx_keeps_typed_values_and_source_row_numbers_of_non_empty_rows() -> None:
    content = xlsx(
        [
            ["Entreprise", "Nom", "Date"],
            ["Transports Exemple SARL", "Test", datetime(2026, 9, 14, 0, 0)],
            [],
            [None, "   ", None],
            ["Fret Modèle", 42, None],
        ]
    )

    workbook = read_workbook("prospects.xlsx", content, LIMITS)

    assert workbook.format is FileFormat.XLSX
    assert workbook.filename == "prospects.xlsx"
    assert len(workbook.fingerprint) == 64
    (sheet,) = workbook.sheets
    assert [row.number for row in sheet.rows] == [1, 2, 5]
    assert sheet.rows[1].values == ("Transports Exemple SARL", "Test", datetime(2026, 9, 14))
    assert sheet.rows[2].values == ("Fret Modèle", 42)  # trailing empty cells trimmed
    assert sheet.width == 3


def test_xlsx_reads_cached_values_never_formulas() -> None:
    content = xlsx([["Entreprise", "Total"], ["Transports Exemple", "=1+1"]])

    (sheet,) = read_workbook("formulas.xlsx", content, LIMITS).sheets

    # A formula without a cached result (never computed by Excel) reads as empty, not as code.
    assert sheet.rows[1].values == ("Transports Exemple",)


def test_xlsx_copies_vertical_merges_down_rows_that_have_content() -> None:
    content = xlsx(
        [
            ["Entreprise", "Nom", "Note"],
            ["Transports Exemple", "Test", "fusion horizontale"],
            [None, "Démo", None],
            [None, None, None],
            [None, "Essai", None],
        ],
        merges=("A2:A5", "C2:D2"),
    )

    (sheet,) = read_workbook("merged.xlsx", content, LIMITS).sheets

    assert sheet.merged_ranges == 2
    assert [row.number for row in sheet.rows] == [1, 2, 3, 5]  # the empty row 4 is not created
    assert [row.value(0) for row in sheet.rows[1:]] == ["Transports Exemple"] * 3
    assert [sorted(row.copied) for row in sheet.rows] == [[], [], [0], [0]]
    assert sheet.rows[1].value(3) is None  # horizontal merges copy nothing


def test_xlsx_reports_every_sheet_in_order() -> None:
    book = OpenpyxlWorkbook()
    first = book.active
    assert first is not None
    first.title = "actualité"
    first.append(["Actualités fictives"])
    book.create_sheet("Base client ").append(["Entreprise", "Nom", "Mail"])
    buffer = io.BytesIO()
    book.save(buffer)

    workbook = read_workbook("two.xlsx", buffer.getvalue(), LIMITS)

    assert [sheet.name for sheet in workbook.sheets] == ["actualité", "Base client "]


@pytest.mark.parametrize(
    ("encoding", "prefix"),
    [("utf-8", b""), ("utf-8", b"\xef\xbb\xbf"), ("utf-16", b"")],
)
def test_csv_unicode_encodings(encoding: str, prefix: bytes) -> None:
    text = f"{CSV_HEADER}\nTransports Exemple;Démo;Hélène;helene.demo@example.com\n"
    content = prefix + text.encode(encoding)

    workbook = read_workbook("export.csv", content, LIMITS)

    assert workbook.format is FileFormat.CSV
    assert workbook.encoding == encoding
    assert workbook.notices == ()
    assert workbook.sheets[0].rows[1].values[2] == "Hélène"


def test_csv_falls_back_to_windows_1252_with_a_notice() -> None:
    content = f"{CSV_HEADER}\nSociété Exemple;Démo;Hélène;h@example.com\n".encode("cp1252")

    workbook = read_workbook("export.csv", content, LIMITS)

    assert workbook.encoding == "cp1252"
    assert [notice.code for notice in workbook.notices] == [DiagnosticCode.FILE_ENCODING_FALLBACK]
    assert workbook.sheets[0].rows[1].values[:3] == ("Société Exemple", "Démo", "Hélène")


def test_csv_undecodable_bytes_are_rejected() -> None:
    assert rejected_code("export.csv", b"Nom;Mail\n\x81\x8d\x90;x\n") is (
        DiagnosticCode.FILE_ENCODING_UNKNOWN
    )


@pytest.mark.parametrize("delimiter", [";", ",", "\t"])
def test_csv_delimiter_is_sniffed(delimiter: str) -> None:
    rows = [
        ["Entreprise", "Nom", "Mail"],
        ["Transports Exemple", "Test", "jean.test@example.com"],
        ["Fret Modèle", "Essai", "paul.essai@example.com"],
    ]
    content = "\n".join(delimiter.join(row) for row in rows).encode()

    workbook = read_workbook("export.csv", content, LIMITS)

    assert workbook.delimiter == delimiter
    assert workbook.sheets[0].rows[2].values == tuple(rows[2])


def test_csv_quoted_fields_keep_other_delimiters_and_line_breaks() -> None:
    content = (
        'Entreprise;Nom;Adresse\n"Exemple, Fils et Cie";Test;"12 rue Fictive\n69000 Lyon"\n'
        "Fret Modèle;Essai;\n"
    ).encode()

    sheet = read_workbook("export.csv", content, LIMITS).sheets[0]

    assert read_workbook("export.csv", content, LIMITS).delimiter == ";"
    assert sheet.rows[1].values == ("Exemple, Fils et Cie", "Test", "12 rue Fictive\n69000 Lyon")
    assert [row.number for row in sheet.rows] == [1, 2, 3]  # records, not physical lines
    assert sheet.rows[2].values == ("Fret Modèle", "Essai")


def test_csv_ambiguous_delimiter_prefers_semicolon() -> None:
    assert read_workbook("one.csv", b"Nom\nTest\n", LIMITS).delimiter == ";"


def test_empty_and_unsupported_files_are_rejected() -> None:
    assert rejected_code("prospects.xlsx", b"") is DiagnosticCode.FILE_EMPTY
    assert rejected_code("prospects.pdf", b"%PDF-1.7") is DiagnosticCode.FILE_UNSUPPORTED_FORMAT
    assert rejected_code("prospects.xls", b"not an OLE file") is (
        DiagnosticCode.FILE_UNSUPPORTED_FORMAT
    )


def test_legacy_xls_and_encrypted_workbooks_are_told_apart() -> None:
    legacy = OLE_MAGIC + b"\x00" * 512
    encrypted = OLE_MAGIC + b"\x00" * 64 + "EncryptionInfo".encode("utf-16-le") + b"\x00" * 64

    assert rejected_code("ancien.xls", legacy) is DiagnosticCode.FILE_LEGACY_XLS
    # Excel saves password-protected .xlsx files as OLE containers too.
    assert rejected_code("protege.xlsx", encrypted) is DiagnosticCode.FILE_ENCRYPTED


@pytest.mark.parametrize(
    "content",
    [
        b"PK\x03\x04 truncated zip",
        b"plain text renamed as xlsx",
    ],
)
def test_corrupt_xlsx_is_rejected_without_echoing_content(content: bytes) -> None:
    with pytest.raises(ImportRejectedError) as caught:
        read_workbook("prospects.xlsx", content, LIMITS)

    assert caught.value.code is DiagnosticCode.FILE_CORRUPT
    assert str(caught.value) == "Le fichier est illisible ou endommagé."
    assert caught.value.__cause__ is None


def test_zip_that_is_not_a_workbook_is_corrupt() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "fichier fictif")

    assert rejected_code("prospects.xlsx", buffer.getvalue()) is DiagnosticCode.FILE_CORRUPT


def test_file_size_limit() -> None:
    limits = ImportLimits(max_file_bytes=1024)

    code = rejected_code("big.csv", b"Nom\n" + b"x" * 1100, limits)

    assert code is DiagnosticCode.FILE_TOO_LARGE


def test_highly_compressed_xlsx_is_refused_before_parsing() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", b"0" * 200_000)
    content = buffer.getvalue()
    limits = ImportLimits(max_file_bytes=max(len(content), 2048))

    assert rejected_code("bomb.xlsx", content, limits) is DiagnosticCode.FILE_TOO_LARGE


def test_row_limit_counts_non_empty_rows() -> None:
    rows = "\n".join(["Nom", "", "A", "", "B", "C"]).encode()

    assert len(read_workbook("rows.csv", rows, ImportLimits(max_rows=4)).sheets[0].rows) == 4
    with pytest.raises(ImportRejectedError) as caught:
        read_workbook("rows.csv", rows, ImportLimits(max_rows=3))
    assert caught.value.code is DiagnosticCode.SHEET_TOO_MANY_ROWS
    assert "3" in str(caught.value)


def test_column_limit_applies_to_xlsx_content() -> None:
    content = xlsx([[f"Col {n}" for n in range(6)]])

    assert rejected_code("wide.xlsx", content, ImportLimits(max_columns=5)) is (
        DiagnosticCode.SHEET_TOO_MANY_COLUMNS
    )


def test_limits_come_from_settings() -> None:
    settings = Settings(import_max_file_mb=2, import_max_rows=10, import_max_columns=7)

    assert ImportLimits.from_settings(settings) == ImportLimits(2 * 1024 * 1024, 10, 7)
