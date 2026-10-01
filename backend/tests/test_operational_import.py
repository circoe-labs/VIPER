"""Import redesign (sequences rework, Slice S2; decisions D3, D4, D5, D10, D11, D12) against the
real test database, on a synthetic workbook in the operational layout
(`tests/fixtures/synthetic/operational_workbook.py`). Every value is invented."""

import io
import uuid
from datetime import date, datetime
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.actor import ActorType
from app.core.business_time import BUSINESS_TIMEZONE
from app.models import (
    Cohort,
    Company,
    ContactMessage,
    ContactSequence,
    ContactTracking,
    Email,
    ImportRowMetadata,
    Prospect,
    QualityAlert,
)
from app.models.enums import (
    ActivityStatus,
    Civility,
    ContactMessageStatus,
    ContactTrackingStatus,
    QualityAlertSource,
    QualityAlertType,
    SendSource,
    VerificationStatus,
)
from app.services import cohorts, companies, contact_sequences, import_commit, prospects
from app.services.contact_tracking import ContactTrackingInput, save_contact_tracking
from app.services.imports.decisions import ImportDecisions, PreviewOptions
from app.services.imports.preview import ImportFile
from app.services.imports.review import DecisionErrorCode, ImportReview
from app.services.imports.text import json_value
from app.services.imports.workbook import ImportLimits
from tests.builders import OPERATOR, bind_operator
from tests.fixtures.synthetic.operational_workbook import (
    JUNK_ROWS,
    ROWS,
    SHEET,
    Row,
    operational_xlsx,
)
from tests.import_support import LEGAL_BASIS, seed

S = ContactTrackingStatus
LIMITS = ImportLimits()
FILENAME = "base_operationnelle.xlsx"
PAST = date(2020, 1, 6)  # the cohorts' real dates: never derived from an ISO week
FUTURE = date(2099, 1, 5)
DATES = {"S37": PAST.isoformat(), "S39": FUTURE.isoformat()}


@pytest.fixture
def ids(db_session: Session) -> dict[str, uuid.UUID]:
    found = seed(db_session)
    bind_operator(db_session)  # from here on, direct writes are a person's (OPERATOR)
    return found


def upload(rows: tuple[Row, ...] | list[Row] = ROWS) -> ImportFile:
    return ImportFile(FILENAME, operational_xlsx(rows))


def review_of(session: Session, file: ImportFile) -> ImportReview:
    review, _ = import_commit.review_upload(session, file, PreviewOptions(), LIMITS)
    return review


def decide(review: ImportReview, **fields: Any) -> ImportDecisions:
    excluded = {
        row.row_number: {"resolution": {"action": "exclude"}}
        for row in review.preview.rows
        if row.row_number in JUNK_ROWS
    }
    fields.setdefault("rows", excluded)
    fields.setdefault("cohort_dates", {k: v for k, v in DATES.items() if needs(review, k)})
    return ImportDecisions.model_validate(
        {
            "file_fingerprint": review.preview.summary.file_fingerprint,
            "preview_digest": review.digest,
            "legal_basis_or_collection_context": LEGAL_BASIS,
            **fields,
        }
    )


def needs(review: ImportReview, code: str) -> bool:
    return any(group.key == code and group.needs_date for group in review.cohorts)


def run(
    session: Session, rows: tuple[Row, ...] | list[Row] = ROWS, **fields: Any
) -> import_commit.CommitResult:
    file = upload(rows)
    review = review_of(session, file)
    return import_commit.commit_import(session, OPERATOR, file, decide(review, **fields), LIMITS)


def by_name(session: Session, last: str) -> Prospect:
    return session.scalars(select(Prospect).where(Prospect.last_name == last)).one()


def sequence_of(session: Session, prospect: Prospect) -> ContactSequence | None:
    return contact_sequences.current_sequence(session, prospect.id)


def state_of(session: Session, prospect: Prospect) -> S | None:
    return session.scalar(
        select(ContactTracking.status).where(ContactTracking.prospect_id == prospect.id)
    )


def sends(session: Session, prospect: Prospect) -> list[ContactMessage]:
    return list(
        session.scalars(select(ContactMessage).where(ContactMessage.prospect_id == prospect.id))
    )


def alerts(session: Session, **where: Any) -> list[QualityAlert]:
    statement = select(QualityAlert).filter_by(**where).order_by(QualityAlert.created_at)
    return list(session.scalars(statement))


def count(session: Session, model: type[Any]) -> int:
    return session.scalar(select(func.count()).select_from(model)) or 0


# --- D10: first sheet only, nothing lost ---------------------------------------------------------


def test_only_the_first_sheet_feeds_the_crm(db_session: Session, ids: dict[str, uuid.UUID]) -> None:
    review = review_of(db_session, upload())

    sheets = {sheet.name: sheet.status.value for sheet in review.preview.summary.sheets}
    assert sheets == {SHEET: "imported", "Feuil1": "skipped", "actualité": "skipped"}
    run(db_session)
    assert not db_session.scalars(select(Prospect).where(Prospect.last_name == "Libre")).all()
    assert not db_session.scalars(
        select(Email).where(Email.address == "marie.libre@example.com")
    ).all()


def test_every_non_empty_cell_of_every_row_is_kept_raw(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    """Imported rows and excluded ones (no identity) keep a snapshot of all their cells."""
    content = operational_xlsx()
    sheet = load_workbook(io.BytesIO(content), read_only=True)[SHEET]
    source = {
        number: {
            cell.column_letter: json_value(cell.value)  # type: ignore[arg-type]
            for cell in row
            if cell.value is not None and str(cell.value).strip() != ""
        }
        for number, row in enumerate(sheet.iter_rows(min_row=2), start=2)
    }

    batch = run(db_session).batch

    traces = {
        trace.source_row_number: trace
        for trace in db_session.scalars(
            select(ImportRowMetadata).where(ImportRowMetadata.import_batch_id == batch.id)
        )
    }
    assert set(traces) == set(source)
    for number, cells in source.items():
        raw = {letter: entry["value"] for letter, entry in traces[number].raw_cells.items()}
        assert cells.items() <= raw.items(), number
        assert traces[number].excluded is (number in JUNK_ROWS)
    assert traces[9].prospect_id is None and traces[9].legacy_metadata == {}
    # The merged `Approche client` is copied down, flagged; `Mode de contact` stays raw only.
    assert traces[3].raw_cells["W"] == {
        "header": "Approche client",
        "value": "Approche fictive",
        "merged": True,
    }
    assert traces[2].legacy_metadata["contact_mode"]["value"] == "Auto"
    assert traces[2].raw_cells["J"] == {"header": "Mode de contact", "value": "Auto"}


# --- D10: `Statut_verification` → activity only (D11), never e-mails -----------------------------


def test_the_status_column_sets_the_activity_never_the_e_mail_verification(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    run(db_session)

    expected = {
        "Alpha": ActivityStatus.ACTIVE,  # Validé
        "Beta": ActivityStatus.ACTIVE,  # validé
        "Gamma": ActivityStatus.INACTIVE,  # Inactif
        "Delta": ActivityStatus.UNKNOWN,  # inconnus
        "Eta": ActivityStatus.UNKNOWN,  # a note
    }
    for last, activity in expected.items():
        person = by_name(db_session, last)
        assert person.activity_status is activity, last
        assert person.employment_verified_at is None, last
        assert {e.verification_status for e in person.emails} <= {VerificationStatus.UNVERIFIED}


def test_a_status_differing_from_a_person_s_activity_is_a_conflict(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    luc = db_session.get(Prospect, ids["luc"])
    assert luc is not None
    luc.activity_status = ActivityStatus.INACTIVE  # a person's decision (OPERATOR is bound)
    db_session.flush()
    again = {
        "status": "Validé",
        "company": "Logistique Démo",
        "last_name": "Exemple",
        "first_name": "Luc",
        "email": "luc.exemple@example.com",
    }

    run(db_session, [again])

    db_session.refresh(luc)
    assert luc.activity_status is ActivityStatus.INACTIVE
    [alert] = alerts(db_session, prospect_id=luc.id)
    assert (alert.type, alert.source) == (
        QualityAlertType.IMPORT_CONFLICT,
        QualityAlertSource.IMPORT,
    )
    assert {
        k: alert.detail[k] for k in ("field", "viper_value", "file_value", "reason", "row")
    } == {
        "field": "activity_status",
        "viper_value": "inactive",
        "file_value": "active",
        "reason": "different",
        "row": 2,
    }


# --- D5: cohorts ---------------------------------------------------------------------------------


def test_a_new_cohort_code_needs_its_real_date(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    file = upload()
    review = review_of(db_session, file)

    groups = [(g.code, g.rows, g.existing, g.needs_date) for g in review.cohorts]
    assert groups == [
        ("S0", [7], True, False),  # seeded: reused
        ("S37", [2], False, True),
        ("S39", [3, 4], False, True),  # `S39` and ` s39 `: one code
    ]
    with pytest.raises(import_commit.InvalidDecisionsError) as caught:
        import_commit.commit_import(
            db_session, OPERATOR, file, decide(review, cohort_dates={"S37": "2020-01-06"}), LIMITS
        )
    assert {(e.code, e.group, e.key) for e in caught.value.errors} == {
        (DecisionErrorCode.COHORT_DATE_REQUIRED, "cohorts", "S39")
    }
    assert count(db_session, Cohort) == 1


def test_an_existing_cohort_is_reused_and_cannot_be_redated_by_an_import(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    s37 = cohorts.create_cohort(db_session, OPERATOR, "S37", date(2026, 9, 7))
    file = upload()
    review = review_of(db_session, file)
    assert next(g for g in review.cohorts if g.code == "S37").cohort_id == s37.id

    with pytest.raises(import_commit.InvalidDecisionsError) as caught:
        import_commit.commit_import(
            db_session, OPERATOR, file, decide(review, cohort_dates=DATES), LIMITS
        )
    assert {(e.code, e.key) for e in caught.value.errors} == {
        (DecisionErrorCode.COHORT_EXISTS, "S37")
    }

    result = run(db_session)

    assert result.counts["cohorts_created"] == 1  # S39 only
    sequence = sequence_of(db_session, by_name(db_session, "Alpha"))
    assert sequence is not None and sequence.cohort_id == s37.id
    assert db_session.get(Cohort, s37.id).starts_on == date(2026, 9, 7)  # type: ignore[union-attr]


def test_cohorts_open_sequences_and_a_past_cohort_records_its_contact(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    result = run(db_session)

    created = {c.code: c for c in db_session.scalars(select(Cohort))}
    assert (created["S37"].starts_on, created["S39"].starts_on) == (PAST, FUTURE)
    assert not created["S37"].needs_review
    alpha = by_name(db_session, "Alpha")
    sequence = sequence_of(db_session, alpha)
    assert sequence is not None and sequence.is_open and sequence.cohort.code == "S37"
    [contact] = sends(db_session, alpha)  # S37 is past: its Contact was sent that day
    assert (contact.rank, contact.status, contact.sent_source) == (
        0,
        ContactMessageStatus.SENT,
        SendSource.IMPORT,
    )
    assert contact.sent_at == datetime(2020, 1, 6, tzinfo=BUSINESS_TIMEZONE)
    for last in ("Beta", "Gamma"):  # S39 is in the future: nothing sent yet
        person = by_name(db_session, last)
        found = sequence_of(db_session, person)
        assert found is not None and found.cohort.code == "S39"
        assert sends(db_session, person) == []
    zeta = by_name(db_session, "Zeta")  # S0: validated, outside the campaign, never contacted
    found = sequence_of(db_session, zeta)
    assert found is not None and found.cohort.code == "S0" and sends(db_session, zeta) == []
    for last in ("Delta", "Epsilon", "Eta"):  # no cohort, file not verified: left as they are
        person = by_name(db_session, last)
        assert sequence_of(db_session, person) is None and state_of(db_session, person) is None
    assert (result.counts["sequences_opened"], result.counts["sends_recorded"]) == (4, 1)
    progress = contact_sequences.prospect_sequence(db_session, alpha.id).progress
    assert progress.sent_count == 1  # the level counts the imported Contact (D1)


def test_a_new_prospect_with_a_cohort_and_an_appointment_keeps_both(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    """The row's appointment does not refuse its cohort (the import gave both): the sequence is
    recorded, nothing is sent nor due under the appointment."""
    row = {"week": "S37", "rdv": "oui", "company": "Atelier Rdv", "last_name": "Rdv"}

    result = run(db_session, [row])

    person = by_name(db_session, "Rdv")
    assert state_of(db_session, person) is S.APPOINTMENT_OBTAINED
    sequence = sequence_of(db_session, person)
    assert sequence is not None and sequence.cohort.code == "S37"
    assert sends(db_session, person) == []
    assert (result.counts["sequences_opened"], result.counts.get("sends_recorded", 0)) == (1, 0)


def test_an_existing_prospect_s_state_keeps_the_file_s_cohort_out(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    save_contact_tracking(
        db_session, OPERATOR, ids["luc"], ContactTrackingInput(S.APPOINTMENT_OBTAINED)
    )
    again = {
        "week": "S37",
        "company": "Logistique Démo",
        "last_name": "Exemple",
        "first_name": "Luc",
    }

    run(db_session, [again])

    luc = db_session.get(Prospect, ids["luc"])
    assert luc is not None
    assert sequence_of(db_session, luc) is None and sends(db_session, luc) == []
    [alert] = alerts(db_session, prospect_id=luc.id)
    assert (alert.detail["field"], alert.detail["viper_value"], alert.detail["file_value"]) == (
        "cohort",
        "appointment_obtained",
        "S37",
    )


def test_a_value_that_is_not_a_cohort_raises_a_data_alert_and_stays_raw(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    review = review_of(db_session, upload())
    assert [(g.text, g.rows) for g in review.not_cohorts] == [("retraité", [6])]

    batch = run(db_session).batch

    epsilon = by_name(db_session, "Epsilon")
    [alert] = alerts(db_session, prospect_id=epsilon.id)
    assert (alert.type, alert.source) == (
        QualityAlertType.DATA_INCONSISTENT,
        QualityAlertSource.IMPORT,
    )
    assert (alert.detail["field"], alert.detail["file_value"]) == ("cohort", "retraité")
    assert alert.detail["import_batch_id"] == str(batch.id)
    trace = db_session.scalars(
        select(ImportRowMetadata).where(ImportRowMetadata.prospect_id == epsilon.id)
    ).one()
    assert trace.legacy_metadata["planned_contact"]["value"] == "retraité"
    assert sequence_of(db_session, epsilon) is None


# --- D4: « Défaillant » only from a file declared verified by a person ---------------------------


def test_rows_without_cohort_become_defaillant_only_in_a_verified_file(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    review = review_of(db_session, upload())
    # Announced before the commit: new prospects without a valid Sxx (blank or `retraité`).
    assert review.disqualified_if_verified == [5, 6, 8, *JUNK_ROWS]

    result = run(db_session, human_verified=True)

    assert result.counts["prospects_disqualified"] == 3
    for last in ("Delta", "Epsilon", "Eta"):
        person = by_name(db_session, last)
        tracking = db_session.scalars(
            select(ContactTracking).where(ContactTracking.prospect_id == person.id)
        ).one()
        assert tracking.status is S.DISQUALIFIED, last
        [history] = tracking.status_history
        # A human decision: the person who validated the import, never the import or the AI.
        assert (history.actor_type, history.actor_id) == (ActorType.HUMAN, OPERATOR.id)
    for last in ("Alpha", "Beta", "Gamma", "Zeta"):  # a cohort (S0 included) is no failure
        assert state_of(db_session, by_name(db_session, last)) in (None, S.NEUTRAL), last


def test_a_verified_file_never_overrides_an_existing_prospect_s_state(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    s39 = cohorts.create_cohort(db_session, OPERATOR, "S39", FUTURE)
    luc = db_session.get(Prospect, ids["luc"])
    assert luc is not None
    contact_sequences.change_cohort(db_session, OPERATOR, luc.id, s39.id)
    again = {"company": "Logistique Démo", "last_name": "Exemple", "first_name": "Luc"}

    run(db_session, [again], human_verified=True)

    assert state_of(db_session, luc) is None
    sequence = sequence_of(db_session, luc)
    assert sequence is not None and sequence.cohort_id == s39.id
    [alert] = alerts(db_session, prospect_id=luc.id)
    assert (alert.detail["field"], alert.detail["viper_value"], alert.detail["file_value"]) == (
        "contact_state",
        "S39",
        "disqualified",
    )


# --- D11: a person's cohort, sequence and fields win ---------------------------------------------


def test_a_different_cohort_keeps_the_person_s_one_with_a_conflict(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    s39 = cohorts.create_cohort(db_session, OPERATOR, "S39", FUTURE)
    luc = db_session.get(Prospect, ids["luc"])
    assert luc is not None
    contact_sequences.change_cohort(db_session, OPERATOR, luc.id, s39.id)
    again = {
        "week": "S37",
        "company": "Logistique Démo",
        "last_name": "Exemple",
        "first_name": "Luc",
    }

    result = run(db_session, [again])

    sequence = sequence_of(db_session, luc)
    assert sequence is not None and sequence.cohort_id == s39.id and sequence.is_open
    assert count(db_session, ContactSequence) == 1 and sends(db_session, luc) == []
    [alert] = alerts(db_session, prospect_id=luc.id)
    assert (alert.detail["field"], alert.detail["viper_value"], alert.detail["file_value"]) == (
        "cohort",
        "S39",
        "S37",
    )
    assert result.counts.get("sequences_opened", 0) == 0
    trace = db_session.scalars(select(ImportRowMetadata)).one()
    assert trace.legacy_metadata["planned_contact"]["value"] == "S37"


def test_an_existing_prospect_without_cohort_is_completed(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    again = {
        "week": "S37",
        "company": "Logistique Démo",
        "last_name": "Exemple",
        "first_name": "Luc",
    }

    run(db_session, [again])

    luc = db_session.get(Prospect, ids["luc"])
    assert luc is not None
    sequence = sequence_of(db_session, luc)
    assert sequence is not None and sequence.cohort.code == "S37"
    assert [m.sent_source for m in sends(db_session, luc)] == [SendSource.IMPORT]
    assert alerts(db_session, prospect_id=luc.id) == []


def test_fields_a_person_set_or_emptied_are_never_changed(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    company = companies.create_company(
        db_session, OPERATOR, companies.CompanyInput(display_name="Atelier Manuel")
    )
    person = prospects.create_prospect(  # created by a person
        db_session,
        OPERATOR,
        prospects.ProspectInput(
            first_name="Rose",
            last_name="Manuelle",
            civility=Civility.MS,
            company_id=company.id,
            exact_job_title="Directrice fictive",
        ),
    )
    person.exact_job_title = None  # then emptied by that person
    db_session.flush()
    row = {
        "company": "Atelier Manuel",
        "last_name": "Manuelle",
        "first_name": "Rose",
        "civility": "M.",  # differs
        "job": "Responsable transport",  # the person emptied it
        "email": "rose.manuelle@example.com",  # empty in VIPER: filled
    }

    run(db_session, [row])

    db_session.refresh(person)
    assert (person.civility, person.exact_job_title) == (Civility.MS, None)
    assert [e.address for e in person.emails] == ["rose.manuelle@example.com"]
    found = {
        (a.detail["field"], a.detail["viper_value"], a.detail["file_value"], a.detail["reason"])
        for a in alerts(db_session, prospect_id=person.id)
    }
    assert found == {
        ("civility", "ms", "mr", "different"),
        ("exact_job_title", None, "Responsable transport", "human_cleared"),
    }


def test_company_fields_follow_the_same_precedence(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    company = companies.create_company(
        db_session,
        OPERATOR,
        companies.CompanyInput(
            display_name="Atelier Témoin",
            project_type="Étude manuelle",
            client_approach="Approche manuelle",
        ),
    )
    stored = db_session.get(Company, company.id)
    assert stored is not None
    stored.client_approach = None  # a person emptied it
    db_session.flush()
    row = {
        "company": "Atelier Témoin",
        "last_name": "Témoin",
        "first_name": "Ugo",
        "project_type": "Étude importée",  # differs
        "approach": "Approche importée",  # emptied by a person
        "references": "Fiche importée",  # empty: filled
    }

    run(db_session, [row])

    db_session.refresh(stored)
    assert (stored.project_type, stored.client_approach, stored.circoe_references) == (
        "Étude manuelle",
        None,
        "Fiche importée",
    )
    found = {
        (a.detail["field"], a.detail["viper_value"], a.detail["file_value"], a.detail["reason"])
        for a in alerts(db_session, company_id=company.id)
    }
    assert found == {
        ("project_type", "Étude manuelle", "Étude importée", "different"),
        ("client_approach", None, "Approche importée", "human_cleared"),
    }


def test_categories_a_person_chose_or_removed_are_kept(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    """A person's creation event lists the categories (an id list): reading it must not break
    the import, whatever the person did with them afterwards."""
    chosen = companies.create_company(
        db_session,
        OPERATOR,
        companies.CompanyInput(
            display_name="Atelier Choisi", activity_category_ids=[ids["storage"]]
        ),
    )
    emptied = companies.create_company(
        db_session,
        OPERATOR,
        companies.CompanyInput(display_name="Atelier Vidé", activity_category_ids=[ids["road"]]),
    )
    stored = db_session.get(Company, emptied.id)
    assert stored is not None
    stored.activity_categories = []  # a person removed them
    db_session.flush()
    category = "Transport routier de marchandises"
    rows = [
        {
            "company": "Atelier Choisi",
            "category": category,
            "last_name": "Choix",
            "first_name": "A",
        },
        {"company": "Atelier Vidé", "category": category, "last_name": "Vide", "first_name": "B"},
    ]

    run(db_session, rows)

    kept = db_session.get(Company, chosen.id)
    assert kept is not None
    db_session.refresh(kept)
    db_session.refresh(stored)
    assert [c.id for c in kept.activity_categories] == [ids["storage"]]
    assert stored.activity_categories == []
    found = {
        (a.company_id, a.detail["field"], a.detail["reason"])
        for a in alerts(db_session, type=QualityAlertType.IMPORT_CONFLICT)
    }
    assert found == {
        (chosen.id, "activity_category_ids", "different"),
        (emptied.id, "activity_category_ids", "human_cleared"),
    }


def test_e_mails_a_person_removed_are_not_brought_back(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    company = companies.create_company(
        db_session, OPERATOR, companies.CompanyInput(display_name="Atelier Iris")
    )
    person = prospects.create_prospect(
        db_session,
        OPERATOR,
        prospects.ProspectInput(
            first_name="Iris",
            last_name="Retirée",
            company_id=company.id,
            emails=[prospects.ChannelInput(value="iris.retiree@example.com", is_primary=True)],
        ),
    )
    stored = db_session.get(Prospect, person.id)
    assert stored is not None
    for email in list(stored.emails):  # a person deleted the address
        db_session.delete(email)
    db_session.flush()
    db_session.expire(stored)
    row = {
        "company": "Atelier Iris",
        "last_name": "Retirée",
        "first_name": "Iris",
        "email": "iris.retiree@example.com",
    }

    run(db_session, [row])

    db_session.refresh(stored)
    assert stored.emails == []
    [alert] = alerts(db_session, prospect_id=person.id)
    assert (alert.detail["field"], alert.detail["reason"]) == ("emails", "human_cleared")


def test_a_referent_set_by_hand_is_left_alone(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    save_contact_tracking(
        db_session, OPERATOR, ids["luc"], ContactTrackingInput(S.NEUTRAL, referent_id=ids["paul"])
    )
    again = {
        "referent": "Claire Référente",
        "company": "Logistique Démo",
        "last_name": "Exemple",
        "first_name": "Luc",
    }

    run(db_session, [again])

    tracking = db_session.scalars(
        select(ContactTracking).where(ContactTracking.prospect_id == ids["luc"])
    ).one()
    assert (tracking.status, tracking.referent_id) == (S.NEUTRAL, ids["paul"])
    [alert] = alerts(db_session, prospect_id=ids["luc"])
    assert (alert.detail["field"], alert.detail["viper_value"]) == ("referent_id", str(ids["paul"]))


# --- idempotence ---------------------------------------------------------------------------------


def test_importing_the_same_file_again_changes_nothing(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    run(db_session, human_verified=True)
    before = {
        model: count(db_session, model)
        for model in (Prospect, Company, Email, Cohort, ContactSequence, ContactMessage)
    }
    before_alerts = count(db_session, QualityAlert)
    before_states = {
        p.last_name: state_of(db_session, p) for p in db_session.scalars(select(Prospect))
    }

    second = run(db_session, human_verified=True, acknowledge_reimport=True)

    after = {model: count(db_session, model) for model in before}
    assert after == before
    assert count(db_session, QualityAlert) == before_alerts
    assert {
        p.last_name: state_of(db_session, p) for p in db_session.scalars(select(Prospect))
    } == before_states
    counts = second.counts
    assert counts.get("prospects_created", 0) == 0 and counts["prospects_attached"] == 7
    assert counts.get("sequences_opened", 0) == counts.get("sends_recorded", 0) == 0
    assert counts.get("prospects_disqualified", 0) == counts["alerts_raised"] == 0


# --- D12: missing values stay empty --------------------------------------------------------------


def test_missing_or_placeholder_values_stay_empty(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    run(db_session)

    delta = by_name(db_session, "Delta")  # job title `0`
    epsilon = by_name(db_session, "Epsilon")  # civility `0`, no e-mail, no phone
    assert delta.exact_job_title is None and delta.role_id is None
    assert epsilon.civility is None
    assert (epsilon.emails, epsilon.phones) == ([], [])


# --- excluded rows -------------------------------------------------------------------------------


def test_an_excluded_opposed_row_is_traced_with_that_prospect(
    db_session: Session, ids: dict[str, uuid.UUID]
) -> None:
    blocked = {
        "week": "S37",
        "company": "Logistique Démo",
        "last_name": "Bloqué",
        "first_name": "Bruno",
        "email": "bruno.bloque@logistique-demo.example",
    }
    other = {"company": "Atelier Libre", "last_name": "Libre", "first_name": "Ana"}

    run(db_session, [blocked, other])

    trace = db_session.scalars(
        select(ImportRowMetadata).where(ImportRowMetadata.source_row_number == 2)
    ).one()
    assert (trace.excluded, trace.prospect_id, trace.company_id) == (True, ids["blocked"], None)
    assert trace.raw_cells["N"] == {"header": "Prénom", "value": "Bruno"}
    blocked_person = db_session.get(Prospect, ids["blocked"])
    assert blocked_person is not None
    assert sequence_of(db_session, blocked_person) is None
