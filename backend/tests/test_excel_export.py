"""Normalized Excel export (Task 10): layout, types, round-trip after manual edits, determinism,
formula guard, audit, protection and scale. Synthetic data only; every value is invented."""

import time as clock
import uuid
from collections.abc import Sequence
from datetime import UTC, datetime
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, insert, select
from sqlalchemy.orm import Session

from app.core.actor import ActorType
from app.models import (
    Company,
    ContactTracking,
    Email,
    Establishment,
    ImportBatch,
    ImportRowMetadata,
    InternalReferent,
    Phone,
    Prospect,
    ProspectSource,
    User,
)
from app.models.enums import (
    ContactTrackingStatus,
    ImportBatchStatus,
    OriginType,
    PhoneType,
    ProspectSourceType,
)
from app.services import excel_export, import_commit, prospects
from app.services.contact_tracking import ContactTrackingInput, save_contact_tracking
from app.services.exports.spec import SHEETS
from app.services.imports.decisions import ImportDecisions, PreviewOptions
from app.services.imports.preview import ImportFile
from app.services.imports.workbook import ImportLimits
from tests.builders import (
    OPERATOR,
    SIREN,
    add_company,
    add_email,
    add_phone,
    add_prospect,
    audit_events,
    bind_operator,
    with_key,
)
from tests.fixtures.synthetic.legacy_workbook import SAMPLE_ROWS, legacy_xlsx
from tests.import_support import FILENAME, LEGAL_BASIS, seed

EXPORT = "/api/exports/workbook"
PARIS = ZoneInfo("Europe/Paris")
GENERATED_AT = datetime(2026, 9, 10, 7, 30, tzinfo=UTC)
LEGACY_STAGE_HEADERS = {"rdv obtenu", "Devis envoyé", "Suivi", "Relance 1", "relance 2"}


def build(session: Session) -> Workbook:
    return load_workbook(
        BytesIO(excel_export.build_workbook(session, generated_at=GENERATED_AT).content)
    )


def records(book: Workbook, sheet: str) -> list[dict[str, Any]]:
    rows = list(book[sheet].iter_rows(values_only=True))
    headers = [str(header) for header in rows[0]]
    return [dict(zip(headers, row, strict=True)) for row in rows[1:]]


def person(rows: Sequence[dict[str, Any]], first: str, last: str) -> dict[str, Any]:
    [row] = [row for row in rows if (row["Prénom"], row["Nom"]) == (first, last)]
    return row


def row_of(rows: Sequence[dict[str, Any]], prospect: Prospect) -> dict[str, Any]:
    [row] = [row for row in rows if row["ID VIPER"] == str(prospect.id)]
    return row


def cell_of(book: Workbook, sheet: str, header: str, row: int = 2) -> Any:
    worksheet = book[sheet]
    column = [cell.value for cell in worksheet[1]].index(header) + 1
    return worksheet.cell(row=row, column=column)


# --- layout -------------------------------------------------------------------------------------


def test_sheets_and_columns_follow_the_specification(db_session: Session) -> None:
    book = build(db_session)

    assert book.sheetnames == [sheet.name for sheet in SHEETS]
    for sheet in SHEETS:
        worksheet = book[sheet.name]
        assert [cell.value for cell in worksheet[1]] == [c.header for c in sheet.columns]
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref is not None and worksheet.auto_filter.ref.startswith("A1:")
        assert worksheet["A1"].font.bold
        first = sheet.columns[0]
        assert worksheet.column_dimensions["A"].width == first.width
    prospect_headers = [column.header for column in SHEETS[0].columns]
    assert prospect_headers[0] == "Référent"  # grill priority: Referent first
    assert not LEGACY_STAGE_HEADERS & set(prospect_headers)  # one status, not five booleans
    assert book.properties.creator == "VIPER"


def test_values_are_typed_dates_and_texts_keep_leading_zeros(db_session: Session) -> None:
    siren = with_key("01234567")  # a SIREN starting with 0
    company = add_company(db_session, "Transports Zéro SARL", siren=siren)
    company.establishments.append(
        Establishment(
            address_line1="1 rue Fictive", postal_code="01000", city="Bourg", is_primary=True
        )
    )
    prospect = add_prospect(
        db_session, company, employment_verified_at=datetime(2026, 9, 1, 22, 30, tzinfo=UTC)
    )
    add_phone(db_session, prospect, "0100000001", is_primary=True)
    save_contact_tracking(
        db_session,
        OPERATOR,
        prospect.id,
        ContactTrackingInput(
            status=ContactTrackingStatus.APPOINTMENT_OBTAINED,
            planned_contact_at=datetime(2026, 1, 1, tzinfo=PARIS),
            appointment_at=datetime(2026, 9, 15, 14, 30, tzinfo=PARIS),
        ),
    )

    book = build(db_session)

    verified = cell_of(book, "Prospects", "Date de vérification")
    # 22:30 UTC on 1 September is already 2 September in Paris.
    assert verified.value == datetime(2026, 9, 2) and verified.number_format == "dd/mm/yyyy"
    appointment = cell_of(book, "Prospects", "Date de rendez-vous")
    assert appointment.value == datetime(2026, 9, 15, 14, 30)
    assert appointment.number_format == "dd/mm/yyyy hh:mm"
    assert cell_of(book, "Prospects", "Semaine").value == "S01 2026"
    for header, value in (("Téléphone", "0100000001"), ("SIREN", siren), ("Code postal", "01000")):
        cell = cell_of(book, "Prospects", header)
        assert (cell.value, cell.data_type, cell.number_format) == (value, "s", "@")
    assert cell_of(book, "Prospects", "ID VIPER").value == str(prospect.id)
    counts = cell_of(book, "Entreprises", "Prospects")
    assert (counts.value, counts.data_type) == (1, "n")


def test_a_prospect_without_company_or_tracking_leaves_those_columns_empty(
    db_session: Session,
) -> None:
    add_prospect(db_session, None, first_name="Solo", last_name="Test")

    [row] = records(build(db_session), "Prospects")

    assert row["Entreprise"] is None and row["ID entreprise"] is None
    assert row["Suivi de contact"] is None and row["Référent"] is None
    assert (row["Statut activité"], row["Ne pas contacter"]) == ("Inconnu", "Non")


def test_rows_are_ordered_by_company_then_person_and_companies_without_prospects_are_kept(
    db_session: Session,
) -> None:
    beta = add_company(db_session, "Beta Fret")
    alpha = add_company(db_session, "alpha Transports")
    add_company(db_session, "Gamma Vide")  # no prospect
    add_prospect(db_session, beta, first_name="Zoé", last_name="Aubert")
    add_prospect(db_session, alpha, first_name="Yves", last_name="Zola")
    add_prospect(db_session, alpha, first_name="Anne", last_name="Élan")
    add_prospect(db_session, None, first_name="Sans", last_name="Entreprise")

    book = build(db_session)

    assert [(r["Entreprise"], r["Nom"]) for r in records(book, "Prospects")] == [
        ("alpha Transports", "Élan"),
        ("alpha Transports", "Zola"),
        ("Beta Fret", "Aubert"),
        (None, "Entreprise"),
    ]
    companies = records(book, "Entreprises")
    assert [(c["Entreprise"], c["Prospects"]) for c in companies] == [
        ("alpha Transports", 2),
        ("Beta Fret", 1),
        ("Gamma Vide", 0),
    ]


# --- round trip ---------------------------------------------------------------------------------


def import_sample(session: Session) -> ImportBatch:
    file = ImportFile(FILENAME, legacy_xlsx(SAMPLE_ROWS))
    limits = ImportLimits()
    review, _ = import_commit.review_upload(session, file, PreviewOptions(), limits)
    decisions = ImportDecisions.model_validate(
        {
            "file_fingerprint": review.preview.summary.file_fingerprint,
            "preview_digest": review.digest,
            "legal_basis_or_collection_context": LEGAL_BASIS,
            "rows": {"10": {"resolution": {"action": "exclude"}}},
            "week_year": 2026,
        }
    )
    return import_commit.commit_import(session, OPERATOR, file, decisions, limits).batch


def imported(session: Session, first: str, last: str) -> Prospect:
    return session.scalars(
        select(Prospect).where(
            Prospect.first_name == first,
            Prospect.last_name == last,
            Prospect.id.in_(select(ImportRowMetadata.prospect_id)),
        )
    ).one()


def test_an_import_edited_by_hand_is_exported_with_the_edits_and_clean_semantics(
    client: TestClient, db_session: Session
) -> None:
    ids = seed(db_session)
    bind_operator(db_session)
    import_sample(db_session)
    jean = imported(db_session, "Jean", "Test")
    marc = imported(db_session, "Marc", "Démo")
    lea = imported(db_session, "Léa", "Modèle")
    paul = imported(db_session, "Paul", "Essai")
    nina = imported(db_session, "Nina", "Homonyme")
    # Manual edits through the domain services.
    prospects.change_company(db_session, OPERATOR, marc.id, ids["demo"])
    prospects.add_channels(
        db_session, OPERATOR, jean, emails=[prospects.ChannelInput("jean.alias@example.com")]
    )
    jean_tracking = jean.contact_tracking
    assert jean_tracking is not None
    save_contact_tracking(
        db_session,
        OPERATOR,
        jean.id,
        ContactTrackingInput(
            status=ContactTrackingStatus.APPOINTMENT_OBTAINED,
            planned_contact_at=jean_tracking.planned_contact_at,
            referent_id=jean_tracking.referent_id,
            response_received_at=datetime(2026, 9, 8, 9, 0, tzinfo=PARIS),
            appointment_at=datetime(2026, 9, 15, 14, 30, tzinfo=PARIS),
        ),
    )
    save_contact_tracking(
        db_session, OPERATOR, lea.id, ContactTrackingInput(ContactTrackingStatus.NOT_INTERESTED)
    )
    save_contact_tracking(
        db_session,
        OPERATOR,
        nina.id,
        ContactTrackingInput(ContactTrackingStatus.TO_CONTACT, referent_id=ids["paul"]),
    )
    prospects.mark_do_not_contact(db_session, OPERATOR, paul.id, reason="Opposition fictive")
    db_session.refresh(jean)
    # Employment verification through the Database Explorer (the manual path of V1 until the
    # Prospect editor, Task 15).
    verified = client.post(
        "/api/explorer/tables/prospects/changes",
        json={
            "updates": [
                {
                    "key": {"id": str(jean.id)},
                    "version": jean.updated_at.isoformat(),
                    "values": {
                        "activity_status": "active",
                        "employment_verified_at": "2026-09-09T10:00:00+02:00",
                    },
                }
            ]
        },
    )
    assert verified.status_code == 200, verified.text

    response = client.get(EXPORT)

    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    people = records(book, "Prospects")
    jean_row = row_of(people, jean)
    assert jean_row["Référent"] == "Claire Référente"
    assert jean_row["Date de contact prévue"] == datetime(2026, 9, 7)  # S37 with the chosen year
    assert jean_row["Semaine"] == "S37 2026"
    assert (jean_row["Entreprise"], jean_row["Catégories d'activité"]) == (
        "Transports Exemple SARL",
        "Transport routier de marchandises",
    )
    assert (jean_row["Civilité"], jean_row["Rôle"], jean_row["Intitulé exact"]) == (
        "M.",
        "Responsable transport",
        "Responsable transport",
    )
    assert (jean_row["Statut activité"], jean_row["Date de vérification"]) == (
        "Actif",
        datetime(2026, 9, 9),
    )
    assert (jean_row["E-mail"], jean_row["Téléphone"], jean_row["Type de téléphone"]) == (
        "jean.test@example.com",
        "+33600000001",
        "Mobile",
    )
    assert jean_row["Suivi de contact"] == "RDV obtenu"
    assert isinstance(jean_row["Statut depuis le"], datetime)
    assert jean_row["Date de réponse"] == datetime(2026, 9, 8)
    assert jean_row["Date de rendez-vous"] == datetime(2026, 9, 15, 14, 30)
    assert (jean_row["Origine"], jean_row["Base légale ou contexte de collecte"]) == (
        "Import Excel",
        LEGAL_BASIS,
    )
    assert jean_row["ID VIPER"] == str(jean.id)
    # Company change: new company, employment verification cleared; `xxx` never a referent.
    marc_row = row_of(people, marc)
    assert (marc_row["Entreprise"], marc_row["Date de vérification"]) == (
        "Logistique Démo SAS",
        None,
    )
    assert marc_row["ID entreprise"] == str(ids["demo"])
    # Referent split: only internal referents, never a legacy marker, note or e-mail.
    referents = {
        f"{r.first_name} {r.last_name}" for r in db_session.scalars(select(InternalReferent))
    }
    assert {row["Référent"] for row in people} == {"Claire Référente", "Paul Démo", None}
    assert {row["Référent"] for row in people} - {None} <= referents
    assert row_of(people, nina)["Référent"] == "Paul Démo"
    # Opposition is distinct from non-interest.
    lea_row = row_of(people, lea)
    assert (lea_row["Suivi de contact"], lea_row["Ne pas contacter"]) == ("Non intéressé", "Non")
    paul_row = row_of(people, paul)
    assert (paul_row["Ne pas contacter"], paul_row["Motif d'opposition"]) == (
        "Oui",
        "Opposition fictive",
    )
    assert (
        isinstance(paul_row["Date d'opposition"], datetime) and paul_row["Suivi de contact"] is None
    )
    assert person(people, "Bruno", "Bloqué")["Ne pas contacter"] == "Oui"
    # Establishment address from the legacy `Adresse` column, on the company.
    emma_row = row_of(people, imported(db_session, "Emma", "Test"))
    assert (emma_row["Adresse"], emma_row["Code postal"], emma_row["Ville"]) == (
        "12 rue de l'Exemple",
        "69000",
        "Lyon",
    )
    assert emma_row["Suivi de contact"] == "RDV obtenu"
    # Aliases: every e-mail and phone, keyed by the prospect id, primary first.
    emails = records(book, "E-mails")
    assert len(emails) == db_session.scalar(select(func.count()).select_from(Email))
    assert [
        (e["E-mail"], e["Principal"], e["Actif"], e["Vérification"], e["Origine"])
        for e in emails
        if e["ID VIPER"] == str(jean.id)
    ] == [
        ("jean.test@example.com", "Oui", "Oui", "Non vérifié", "Import"),
        ("jean.alias@example.com", "Non", "Oui", "Non vérifié", "Saisie manuelle"),
    ]
    phones = records(book, "Téléphones")
    assert len(phones) == db_session.scalar(select(func.count()).select_from(Phone))
    assert [
        (p["Numéro"], p["Type"], p["Principal"]) for p in phones if p["ID VIPER"] == str(jean.id)
    ] == [
        ("+33600000001", "Mobile", "Oui"),
        ("+33100000001", "Fixe", "Non"),
    ]
    assert {p["ID VIPER"] for p in phones} <= {row["ID VIPER"] for row in people}
    # Legacy metadata: every preserved value of every imported row, nothing more.
    legacy = records(book, "Données d'origine")
    traces = db_session.scalars(select(ImportRowMetadata)).all()
    assert len(legacy) == sum(len(trace.legacy_metadata) for trace in traces)
    jean_legacy = {
        (x["En-tête d'origine"], x["Valeur d'origine"])
        for x in legacy
        if x["ID VIPER"] == str(jean.id)
    }
    assert {("Mode de contact", "Auto"), ("A contacter", "Oui")} <= jean_legacy
    marc_legacy = {x["Champ"]: x for x in legacy if x["ID VIPER"] == str(marc.id)}
    assert (marc_legacy["referent"]["Valeur d'origine"], marc_legacy["referent"]["Ligne"]) == (
        "xxx",
        3,
    )
    assert marc_legacy["referent"]["Fichier importé"] == FILENAME
    emma_legacy = [x for x in legacy if x["Champ"] == "column_X"]
    assert [(x["Valeur d'origine"], x["Motif"]) for x in emma_legacy] == [
        ("note fictive", "Colonne non reconnue")
    ]
    # Provenance: one source per imported row.
    sources = records(book, "Provenance")
    assert len(sources) == db_session.scalar(select(func.count()).select_from(ProspectSource))
    assert {s["Origine"] for s in sources} == {"Import Excel"}
    assert len(records(book, "Établissements")) == db_session.scalar(
        select(func.count()).select_from(Establishment)
    )


def test_two_exports_of_the_same_data_are_identical(db_session: Session) -> None:
    seed(db_session)
    bind_operator(db_session)
    import_sample(db_session)

    first = excel_export.build_workbook(db_session, generated_at=GENERATED_AT)
    second = excel_export.build_workbook(db_session, generated_at=GENERATED_AT)

    assert first.content == second.content
    assert first.filename == "VIPER_export_2026-09-10.xlsx"


def test_texts_are_never_formulas_and_keep_their_exact_value(db_session: Session) -> None:
    dangerous = ['=HYPERLINK("http://example.com")', "@SUM(1)", "-2+3", "+cmd", "#N/A"]
    company = add_company(db_session, dangerous[0])
    prospect = add_prospect(db_session, company, first_name=dangerous[1], last_name=dangerous[2])
    prospect.exact_job_title = dangerous[3] + "\x07"  # a control character XML cannot hold
    prospect.do_not_contact_reason = None
    add_phone(db_session, prospect, "+33600000009", is_primary=True, type=PhoneType.MOBILE)
    add_email(db_session, prospect, "formule@example.com", is_primary=True)
    company.client_approach = dangerous[4]
    db_session.flush()

    book = build(db_session)

    cells = [cell for worksheet in book.worksheets for row in worksheet.iter_rows() for cell in row]
    assert not [cell for cell in cells if cell.data_type == "f"]
    row = records(book, "Prospects")[0]
    assert (row["Entreprise"], row["Prénom"], row["Nom"]) == tuple(dangerous[:3])
    assert (row["Intitulé exact"], row["Approche client"]) == (dangerous[3], dangerous[4])
    for header in ("Entreprise", "Prénom", "Nom", "Intitulé exact"):
        assert cell_of(book, "Prospects", header).quotePrefix
    phone = cell_of(book, "Prospects", "Téléphone")
    assert (phone.value, phone.data_type, phone.quotePrefix) == ("+33600000009", "s", False)
    assert cell_of(book, "Prospects", "Approche client").data_type == "s"


# --- API ----------------------------------------------------------------------------------------


def test_the_download_is_an_audited_xlsx_attachment(
    client: TestClient, db_session: Session, pilot_user: User
) -> None:
    add_prospect(db_session, add_company(db_session, siren=SIREN))

    response = client.get(EXPORT)

    assert response.status_code == 200
    assert response.headers["content-type"] == excel_export.MEDIA_TYPE
    stamp = datetime.now(PARIS).strftime("%Y-%m-%d")
    assert response.headers["content-disposition"] == (
        f'attachment; filename="VIPER_export_{stamp}.xlsx"'
    )
    assert response.headers["cache-control"] == "no-store"
    assert response.content[:4] == b"PK\x03\x04"
    [event] = audit_events(db_session, action="export.generated")
    assert (event.actor_type, event.actor_id) == (ActorType.HUMAN, str(pilot_user.id))
    assert (event.entity_type, event.entity_id) == (excel_export.EXPORT_ENTITY, None)
    assert event.context["source"] == "ui"
    assert event.changes["prospects_rows"] == {"before": None, "after": 1}
    assert event.changes["companies_rows"]["after"] == 1
    assert event.changes["size_bytes"]["after"] == len(response.content)
    assert set(event.changes) == {f"{sheet.key}_rows" for sheet in SHEETS} | {"size_bytes"}


def test_the_export_requires_a_session(anonymous_client: TestClient) -> None:
    assert anonymous_client.get(EXPORT).status_code == 401


# --- scale --------------------------------------------------------------------------------------


def test_a_few_thousand_prospects_export_in_reasonable_time(db_session: Session) -> None:
    companies = [
        {"id": uuid.uuid7(), "display_name": f"Transports Volume {n:04d}"} for n in range(600)
    ]
    people = [
        {
            "id": uuid.uuid7(),
            "company_id": companies[n % 600]["id"],
            "first_name": f"Prénom{n}",
            "last_name": f"Volume{n:05d}",
        }
        for n in range(3000)
    ]
    batch_id = uuid.uuid7()
    db_session.execute(insert(Company), companies)
    db_session.execute(insert(Prospect), people)
    db_session.execute(
        insert(Email),
        [
            {
                "prospect_id": p["id"],
                "address": f"volume{n}@example.com",
                "is_primary": True,
                "origin_type": OriginType.IMPORTED,
            }
            for n, p in enumerate(people)
        ],
    )
    db_session.execute(
        insert(Phone),
        [
            {
                "prospect_id": p["id"],
                "number": f"+3360000{n:04d}",
                "type": PhoneType.MOBILE,
                "is_primary": True,
                "origin_type": OriginType.IMPORTED,
            }
            for n, p in enumerate(people)
        ],
    )
    db_session.execute(
        insert(ContactTracking),
        [{"prospect_id": p["id"], "status": ContactTrackingStatus.CONTACTED} for p in people[::2]],
    )
    db_session.execute(
        insert(ImportBatch),
        [
            {
                "id": batch_id,
                "filename": "volume.xlsx",
                "status": ImportBatchStatus.PENDING,
                "actor_type": ActorType.HUMAN,
                "actor_display": "Opératrice Test",
            }
        ],
    )
    db_session.execute(
        insert(ProspectSource),
        [
            {
                "prospect_id": p["id"],
                "source_type": ProspectSourceType.EXCEL_IMPORT,
                "import_batch_id": batch_id,
            }
            for p in people
        ],
    )
    db_session.execute(
        insert(ImportRowMetadata),
        [
            {
                "import_batch_id": batch_id,
                "source_sheet": "Base",
                "source_row_number": n + 2,
                "prospect_id": p["id"],
                "legacy_metadata": {
                    "contact_mode": {
                        "column": "I",
                        "header": "Mode de contact",
                        "value": "Auto",
                        "reason": "opaque_field",
                    },
                    "column_X": {
                        "column": "X",
                        "header": None,
                        "value": n,
                        "reason": "unmapped_column",
                    },
                },
            }
            for n, p in enumerate(people)
        ],
    )

    started = clock.perf_counter()
    exported = excel_export.build_workbook(db_session, generated_at=GENERATED_AT)
    elapsed = clock.perf_counter() - started

    assert exported.rows == {
        "prospects": 3000,
        "companies": 600,
        "establishments": 0,
        "emails": 3000,
        "phones": 3000,
        "sources": 3000,
        "legacy_values": 6000,
    }
    assert elapsed < 45, f"export took {elapsed:.1f} s"
    print(f"\n3000 prospects exported in {elapsed:.2f} s, {len(exported.content)} bytes")
