"""Synthetic workbooks in the historical `BASE_CLIENT` layout, generated in memory with openpyxl.

Every value is invented (`example.com` addresses, `+33 6 00 00 00 0x` numbers, fictitious people
and companies). The layout — 24 columns with the repeated `A contacter` header and a trailing
unnamed column, a `Base client ` sheet with a trailing space, an extra `actualité` sheet — mirrors
the structure described in `doc/legacy/legacy-data-profile.md`, never its content.
"""

import io
import uuid
from collections.abc import Mapping, Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models.enums import ContactabilityStatus
from app.services.imports.reference import (
    ImportReferenceData,
    ReferenceCompany,
    ReferenceProspect,
    ReferenceReferent,
    ReferenceTaxonomy,
)

PROSPECT_SHEET = "Base client "
HEADERS: tuple[str | None, ...] = (
    "Référent",
    "A contacter ",
    "Entreprise",
    "rdv obtenu",
    "Devis envoyé",
    "Suivi",
    "Relance 1",
    "relance 2",
    "Mode de contact",
    "Catégorie",
    "Civilité ",
    "Nom",
    "Prénom",
    "Fonction",
    "Mail",
    "Téléphone",
    "Mobile",
    "Adresse ",
    "Projet déjà réalisé avec l'entreprise",
    "Type de projet",
    "Liste des fiches projets_references_CIRCOE.csv",
    "Approche client",
    "A contacter",
    None,
)
# Short keys for the 24 columns, in order.
KEYS = (
    "referent",
    "week",
    "company",
    "rdv",
    "devis",
    "suivi",
    "relance1",
    "relance2",
    "mode",
    "category",
    "civility",
    "last_name",
    "first_name",
    "job",
    "email",
    "phone",
    "mobile",
    "address",
    "project_done",
    "project_type",
    "references",
    "approach",
    "flag",
    "unnamed",
)
COLUMN = {key: index for index, key in enumerate(KEYS)}
NEWS_ROWS: tuple[tuple[Any, ...], ...] = (
    ("Actualités fictives",),
    ("2026-01-05", "Salon fictif du transport"),
    ("2026-02-10", "Webinaire de démonstration"),
)

type Row = Mapping[str, Any]


def legacy_values(row: Row) -> list[Any]:
    unknown = set(row) - set(KEYS)
    assert not unknown, f"unknown fixture keys: {sorted(unknown)}"
    return [row.get(key) for key in KEYS]


def legacy_xlsx(
    rows: Sequence[Row],
    *,
    sheet: str = PROSPECT_SHEET,
    headers: Sequence[str | None] = HEADERS,
    extra_sheets: Mapping[str, Sequence[Sequence[Any]]] | None = None,
    merges: Sequence[str] = (),
    blank_rows_before: int = 0,
) -> bytes:
    """An XLSX with the legacy header row (bold, so the unnamed 24th column is part of the used
    range like in the real file), `rows`, then `extra_sheets` (default: an `actualité` sheet)."""
    book = Workbook()
    worksheet = book.active
    assert worksheet is not None
    worksheet.title = sheet
    for _ in range(blank_rows_before):
        worksheet.append([])
    worksheet.append(list(headers))
    for cell in worksheet[worksheet.max_row]:
        cell.font = Font(bold=True)
    for row in rows:
        worksheet.append(legacy_values(row) if isinstance(row, Mapping) else list(row))
    for cell_range in merges:
        worksheet.merge_cells(cell_range)
    for name, data in (
        extra_sheets if extra_sheets is not None else {"actualité": NEWS_ROWS}
    ).items():
        other = book.create_sheet(name)
        for values in data:
            other.append(list(values))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def ids(*names: str) -> dict[str, uuid.UUID]:
    """Stable fake ids, so previews built from the same reference data compare equal."""
    return {name: uuid.uuid5(uuid.NAMESPACE_URL, f"viper-test:{name}") for name in names}


ID = ids(
    "role-transport",
    "role-dirigeant",
    "role-old",
    "category-road",
    "category-storage",
    "category-old",
    "segment-carrier",
    "referent-claire",
    "referent-paul",
    "referent-former",
    "company-demo",
    "company-sample",
    "prospect-blocked",
    "prospect-luc",
    "prospect-homonym",
)

REFERENCE = ImportReferenceData(
    roles=(
        ReferenceTaxonomy(ID["role-dirigeant"], "Dirigeant", "dirigeant"),
        ReferenceTaxonomy(ID["role-transport"], "Responsable transport", "responsable-transport"),
        ReferenceTaxonomy(ID["role-old"], "Chef de quai", "chef-de-quai", active=False),
    ),
    activity_categories=(
        ReferenceTaxonomy(
            ID["category-road"],
            "Transport routier de marchandises",
            "transport-routier-marchandises",
        ),
        ReferenceTaxonomy(
            ID["category-storage"], "Entreposage et stockage", "entreposage-stockage"
        ),
        ReferenceTaxonomy(ID["category-old"], "Déménagement", "demenagement", active=False),
    ),
    commercial_segments=(ReferenceTaxonomy(ID["segment-carrier"], "Transporteur", "transporteur"),),
    referents=(
        ReferenceReferent(
            ID["referent-claire"], "Claire", "Référente", "claire.referente@example.com"
        ),
        ReferenceReferent(ID["referent-paul"], "Paul", "Démo"),
        ReferenceReferent(ID["referent-former"], "Hugo", "Ancien", active=False),
    ),
    companies=(
        ReferenceCompany(
            ID["company-demo"], "Logistique Démo SAS", email_domain="logistique-demo.example"
        ),
        ReferenceCompany(ID["company-sample"], "Transports Échantillon"),
    ),
    prospects=(
        ReferenceProspect(
            ID["prospect-blocked"],
            "Bruno",
            "Bloqué",
            ID["company-demo"],
            ContactabilityStatus.DO_NOT_CONTACT,
            ("bruno.bloque@logistique-demo.example",),
        ),
        ReferenceProspect(
            ID["prospect-luc"],
            "Luc",
            "Exemple",
            ID["company-demo"],
            emails=("luc.exemple@example.com",),
        ),
        ReferenceProspect(
            ID["prospect-homonym"],
            "Nina",
            "Homonyme",
            None,
            ContactabilityStatus.DO_NOT_CONTACT,
        ),
    ),
)

# One row per legacy compatibility case (doc/process/testing-strategy.md); row numbers in the
# sheet are the list index + 2 (row 1 is the header).
SAMPLE_ROWS: tuple[Row, ...] = (
    {  # 2: clean row — exact role, category, referent; week without year; opaque columns
        "referent": "Claire Référente",
        "week": "S37",
        "company": "Transports Exemple SARL",
        "mode": "Auto",
        "category": "Transport routier de marchandises",
        "civility": "M.",
        "last_name": "TEST",
        "first_name": "Jean",
        "job": "Responsable transport",
        "email": "jean.test@example.com",
        "phone": "01 00 00 00 01",
        "mobile": "06 00 00 00 01",
        "project_done": "Oui",
        "project_type": "1 Étude fictive",
        "flag": "Oui",
    },
    {  # 3: MR, xxx marker, s39, invalid category, company variant, duplicate email, unknown role
        "referent": "xxx",
        "week": "s39",
        "company": "TRANSPORTS EXEMPLE",
        "mode": "Commercial",
        "category": "Non",
        "civility": "MR",
        "last_name": "Démo",
        "first_name": "Marc",
        "job": "Directeur des opérations fictives",
        "email": "JEAN.TEST@example.com",
        "project_done": "Non",
        "project_type": 0,
        "flag": "OUI",
    },
    {  # 4: MME., `?` marker, retraité, two emails and two phones in one cell each
        "referent": "?",
        "week": "retraité",
        "company": "Messagerie Fictive",
        "mode": "Commerciale",
        "category": "2. Logistique & Stockage / Transporteur",
        "civility": "MME.",
        "last_name": "Exemple",
        "first_name": "Claire",
        "email": "claire.exemple@example.com / c.exemple@example.com",
        "phone": "01 00 00 00 03 / 06 00 00 00 02",
    },
    {  # 5: `0` civility, `v` marker, number-typed phone that lost its leading zero
        "referent": "v",
        "company": "Fret Modèle SAS",
        "civility": 0,
        "last_name": "Modèle",
        "first_name": "Léa",
        "mobile": 600000004,
        "email": "lea.modele@example.com",
    },
    {  # 6: email-like referent, invalid email, invalid phone
        "referent": "contact.test@example.com",
        "company": "Fret Modèle",
        "civility": "Mme",
        "last_name": "Essai",
        "first_name": "Paul",
        "email": "paul.essai@",
        "phone": "01 00 00",
    },
    {  # 7: referent note with a departure hint; historical stage conflict
        "referent": "parti à la retraite en 2024",
        "company": "Affrètement Fictif",
        "civility": "M",
        "last_name": "Fictif",
        "first_name": "Hugo",
        "rdv": "non",
        "devis": "oui",
        "relance2": "x",
    },
    {  # 8: consistent stages; unnamed column with a value; address
        "company": "Entrepôts Test",
        "civility": "Mme",
        "last_name": "Test",
        "first_name": "Emma",
        "relance1": "x",
        "rdv": "oui",
        "address": "12 rue de l'Exemple\n69000 Lyon",
        "unnamed": "note fictive",
    },
    {  # 9: matches a do-not-contact prospect by email
        "company": "Logistique Démo",
        "civility": "M.",
        "last_name": "Bloqué",
        "first_name": "Bruno",
        "email": "bruno.bloque@logistique-demo.example",
    },
    {  # 10: no name at all (company-only row)
        "company": "Transports Exemple SARL",
        "category": "Transport routier de marchandises",
    },
    {  # 11: no company; homonym of a do-not-contact prospect; partial referent; inactive role
        "referent": "Paul",
        "civility": "Mme",
        "last_name": "Homonyme",
        "first_name": "Nina",
        "job": "Chef de quai",
    },
)
