"""Synthetic workbook in the layout of the current operational prospecting file (Slice S2).

Structure only, every value invented (`example.com` addresses, `+33 6 00 00 00 xx` numbers,
fictitious people and companies): a first sheet `Base client ` (trailing space) with one header
row — `Statut_verification` first, then the 24 historical columns (the cohort column `A contacter `,
the hidden references column, the second `A contacter` flag and an unnamed empty column) —, a
merged cell, rows without any identity (fragments in the first two columns only), and two more
sheets (`Feuil1` without header, `actualité` free notes) that must never feed the CRM.
"""

import io
from collections.abc import Mapping, Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from tests.fixtures.synthetic.legacy_workbook import HEADERS, KEYS, NEWS_ROWS

SHEET = "Base client "
OPERATIONAL_HEADERS: tuple[str | None, ...] = ("Statut_verification", *HEADERS)
OPERATIONAL_KEYS: tuple[str, ...] = ("status", *KEYS)
LETTER = {key: chr(ord("A") + index) for index, key in enumerate(OPERATIONAL_KEYS)}
# `Feuil1`: prospect-looking values without any header — never imported.
LOOSE_ROWS: tuple[tuple[Any, ...], ...] = (
    ("Transports Feuille Libre", "Marie", "Libre", "marie.libre@example.com"),
    (None, "fragment fictif"),
)

type Row = Mapping[str, Any]

# Rows in sheet order (sheet row = index + 2). Cohort codes: `s37`/`S39` written in several ways,
# `S0`, blank, and a word that is no cohort; `Statut_verification` in its usual spellings and a
# note; `Mode de contact` raw only.
ROWS: tuple[Row, ...] = (
    {  # 2
        "status": "Validé",
        "referent": "Claire Référente",
        "week": "s37",
        "company": "Transports Cohorte SARL",
        "mode": "Auto",
        "category": "Transport routier de marchandises",
        "civility": "M.",
        "last_name": "Alpha",
        "first_name": "Jean",
        "job": "Responsable transport",
        "email": "jean.alpha@example.com",
        "mobile": "06 00 00 00 21",
        "approach": "Approche fictive",
        "flag": "Oui",
    },
    {  # 3
        "status": "validé",
        "week": "S39",
        "company": "Transports Cohorte",
        "mode": "Commercial",
        "civility": "Mme",
        "last_name": "Beta",
        "first_name": "Anne",
        "job": "Dirigeant",
        "email": "anne.beta@example.com",
        "approach": "Approche fictive",
    },
    {  # 4
        "status": "Inactif",
        "week": " s39 ",
        "company": "Messagerie Cohorte",
        "mode": "Auto",
        "civility": "M",
        "last_name": "Gamma",
        "first_name": "Paul",
        "email": "paul.gamma@example.com",
        "references": "Fiche fictive 12",
    },
    {  # 5: no cohort, « Inconnus »
        "status": "inconnus",
        "company": "Fret Cohorte",
        "civility": "Mme",
        "last_name": "Delta",
        "first_name": "Léa",
        "job": 0,  # placeholder: no job title
        "email": "lea.delta@example.com",
    },
    {  # 6: not a cohort
        "status": "Inconnus",
        "week": "retraité",
        "company": "Fret Cohorte",
        "civility": 0,
        "last_name": "Epsilon",
        "first_name": "Marc",
    },
    {  # 7: S0, validated outside the campaign
        "week": "S0",
        "company": "Logistique Hors Campagne",
        "civility": "M.",
        "last_name": "Zeta",
        "first_name": "Hugo",
        "email": "hugo.zeta@example.com",
        "phone": "01 00 00 00 27",
    },
    {  # 8: a note in the status column, no cohort, no channel
        "status": "à revoir",
        "company": "Entrepôts Cohorte",
        "mode": "Commercial",
        "civility": "Mme",
        "last_name": "Eta",
        "first_name": "Nina",
    },
    # 9-14: rows without any identity (no company, name or e-mail): fragments only.
    {"status": "?"},
    {"referent": "xxx"},
    {"status": "v", "referent": "v"},
    {"referent": "note fictive"},
    {"status": "Validé"},
    {"referent": "?", "status": "inactif"},
)
JUNK_ROWS = tuple(range(9, 15))


def operational_values(row: Row) -> list[Any]:
    unknown = set(row) - set(OPERATIONAL_KEYS)
    assert not unknown, f"unknown fixture keys: {sorted(unknown)}"
    return [row.get(key) for key in OPERATIONAL_KEYS]


def operational_xlsx(
    rows: Sequence[Row] = ROWS,
    *,
    merges: Sequence[str] = ("W2:W3",),  # `Approche client` of rows 2-3, merged vertically
    extra_sheets: Mapping[str, Sequence[Sequence[Any]]] | None = None,
) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = SHEET
    sheet.append(list(OPERATIONAL_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)  # the unnamed last column is part of the used range
    for row in rows:
        sheet.append(operational_values(row))
    for cell_range in merges:
        sheet.merge_cells(cell_range)
    sheet.column_dimensions[LETTER["references"]].hidden = True
    others = (
        extra_sheets if extra_sheets is not None else {"Feuil1": LOOSE_ROWS, "actualité": NEWS_ROWS}
    )
    for name, data in others.items():
        other = book.create_sheet(name)
        for values in data:
            other.append(list(values))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
