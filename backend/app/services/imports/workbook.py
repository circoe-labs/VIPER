"""Spreadsheet adapters: uploaded bytes + file name → typed, bounded workbook model.

No legacy knowledge here (that is `fields.py`/`layout.py`). XLSX is read with openpyxl in read-only
mode, cached values only (never formulas); CSV is decoded UTF-8 (BOM allowed) with a Windows-1252
fallback and its delimiter sniffed among `;`, `,` and tab. Only non-empty rows are kept, with their
source row numbers. Failures raise `ImportRejectedError` with a catalogue code and never echo file
content (library exceptions are not chained, as their messages may quote cell values).
"""

import csv
import hashlib
import io
import posixpath
import zipfile
from collections import Counter
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time, timedelta
from enum import StrEnum
from itertools import islice
from pathlib import PurePath
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from app.core.config import Settings
from app.services.imports.diagnostics import Diagnostic, DiagnosticCode, ImportRejectedError
from app.services.imports.diagnostics import diagnostic as make_diagnostic
from app.services.imports.text import CellValue, is_blank

XLSX_SUFFIXES = frozenset({".xlsx", ".xlsm"})
CSV_SUFFIXES = frozenset({".csv", ".tsv", ".txt"})
OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")  # legacy .xls and password-protected workbooks
ENCRYPTION_STREAM = "EncryptionInfo".encode("utf-16-le")
# Cap on the unzipped size of an XLSX, as a multiple of the file-size limit (zip-bomb guard:
# openpyxl loads shared strings in memory even in read-only mode).
UNCOMPRESSED_RATIO = 10
CSV_DELIMITERS = (";", ",", "\t")  # sniffing tie-break order: French Excel exports use `;`
SNIFF_RECORDS = 50
MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
RELATIONSHIP_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PACKAGE_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


class FileFormat(StrEnum):
    XLSX = "xlsx"
    CSV = "csv"


@dataclass(frozen=True, slots=True)
class ImportLimits:
    max_file_bytes: int = 10 * 1024 * 1024
    max_rows: int = 5000  # non-empty rows per sheet
    max_columns: int = 100

    @classmethod
    def from_settings(cls, settings: Settings) -> ImportLimits:
        return cls(
            max_file_bytes=settings.import_max_file_mb * 1024 * 1024,
            max_rows=settings.import_max_rows,
            max_columns=settings.import_max_columns,
        )


@dataclass(frozen=True, slots=True)
class SheetRow:
    number: int  # 1-based row number in the source
    values: tuple[CellValue, ...]
    copied: frozenset[int] = frozenset()  # column indexes filled from a vertical merge

    def value(self, index: int) -> CellValue:
        return self.values[index] if index < len(self.values) else None


@dataclass(frozen=True, slots=True)
class Sheet:
    name: str
    rows: tuple[SheetRow, ...] = ()  # non-empty rows only, in source order
    width: int = 0  # column count: content, or the declared used range when within limits
    merged_ranges: int = 0


@dataclass(frozen=True, slots=True)
class Workbook:
    filename: str
    format: FileFormat
    fingerprint: str  # SHA-256 of the uploaded bytes
    sheets: tuple[Sheet, ...]
    encoding: str | None = None
    delimiter: str | None = None
    notices: tuple[Diagnostic, ...] = field(default=())


def read_workbook(filename: str, content: bytes, limits: ImportLimits) -> Workbook:
    if not content:
        raise ImportRejectedError(DiagnosticCode.FILE_EMPTY)
    if len(content) > limits.max_file_bytes:
        raise too_large(limits)
    if content.startswith(OLE_MAGIC):
        encrypted = ENCRYPTION_STREAM in content
        raise ImportRejectedError(
            DiagnosticCode.FILE_ENCRYPTED if encrypted else DiagnosticCode.FILE_LEGACY_XLS
        )
    suffix = PurePath(filename).suffix.lower()
    fingerprint = hashlib.sha256(content).hexdigest()
    name = PurePath(filename).name
    if suffix in XLSX_SUFFIXES:
        return Workbook(name, FileFormat.XLSX, fingerprint, read_xlsx(content, limits))
    if suffix in CSV_SUFFIXES:
        text, encoding = decode_csv(content)
        delimiter = sniff_delimiter(text)
        notices = (
            (make_diagnostic(DiagnosticCode.FILE_ENCODING_FALLBACK),)
            if encoding == "cp1252"
            else ()
        )
        sheet = read_csv(name, text, delimiter, limits)
        return Workbook(name, FileFormat.CSV, fingerprint, (sheet,), encoding, delimiter, notices)
    raise ImportRejectedError(DiagnosticCode.FILE_UNSUPPORTED_FORMAT)


def too_large(limits: ImportLimits) -> ImportRejectedError:
    return ImportRejectedError(
        DiagnosticCode.FILE_TOO_LARGE, limit_mb=limits.max_file_bytes // (1024 * 1024)
    )


# --- XLSX ---------------------------------------------------------------------------------------


def read_xlsx(content: bytes, limits: ImportLimits) -> tuple[Sheet, ...]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > (
                limits.max_file_bytes * UNCOMPRESSED_RATIO
            ):
                raise too_large(limits)
            merges = merged_ranges(archive)
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ImportRejectedError:
        raise
    except Exception:  # untrusted input: any parse failure means an unreadable file
        raise ImportRejectedError(DiagnosticCode.FILE_CORRUPT) from None
    try:
        sheets = []
        for name in book.sheetnames:
            worksheet = book[name]
            if isinstance(worksheet, ReadOnlyWorksheet):
                sheet = read_worksheet(name, worksheet, limits)
                sheets.append(apply_merges(sheet, merges.get(name, ())))
            else:  # chart sheet: no cells, still reported (and skipped) by name
                sheets.append(Sheet(name))
        return tuple(sheets)
    except ImportRejectedError:
        raise
    except Exception:
        raise ImportRejectedError(DiagnosticCode.FILE_CORRUPT) from None
    finally:
        book.close()


def read_worksheet(name: str, worksheet: ReadOnlyWorksheet, limits: ImportLimits) -> Sheet:
    declared_width = worksheet.max_column or 0
    # Ignore the declared used range while reading: a formatted column XFD would otherwise pad
    # every row to 16 384 cells.
    worksheet.reset_dimensions()
    rows: list[SheetRow] = []
    for number, values in enumerate(worksheet.values, start=1):
        cells = trimmed(tuple(cell_value(value) for value in values))
        if cells:
            check_bounds(name, len(rows) + 1, len(cells), limits)
            rows.append(SheetRow(number, cells))
    content_width = max((len(row.values) for row in rows), default=0)
    width = declared_width if content_width < declared_width <= limits.max_columns else 0
    return Sheet(name, tuple(rows), max(content_width, width))


def cell_value(value: object) -> CellValue:
    if value is None or isinstance(
        value, str | int | float | bool | datetime | date | time | timedelta
    ):
        return value
    return str(value)


def trimmed(values: tuple[CellValue, ...]) -> tuple[CellValue, ...]:
    end = len(values)
    while end and is_blank(values[end - 1]):
        end -= 1
    return values[:end]


def check_bounds(sheet: str, rows: int, columns: int, limits: ImportLimits) -> None:
    if rows > limits.max_rows:
        raise ImportRejectedError(
            DiagnosticCode.SHEET_TOO_MANY_ROWS, sheet=sheet, limit=limits.max_rows
        )
    if columns > limits.max_columns:
        raise ImportRejectedError(
            DiagnosticCode.SHEET_TOO_MANY_COLUMNS, sheet=sheet, limit=limits.max_columns
        )


type Bounds = tuple[int, int, int, int]  # min column, min row, max column, max row (1-based)


def merged_ranges(archive: zipfile.ZipFile) -> dict[str, tuple[Bounds, ...]]:
    """Merged ranges per sheet name. openpyxl's read-only mode does not expose them, so the
    sheet parts are located through the workbook relationships and scanned with the stdlib."""
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        rel.get("Id"): rel.get("Target", "")
        for rel in relationships.iter(f"{PACKAGE_NS}Relationship")
    }
    result: dict[str, tuple[Bounds, ...]] = {}
    for sheet in workbook.iter(f"{MAIN_NS}sheet"):
        target = targets.get(sheet.get(f"{RELATIONSHIP_NS}id"))
        if not target:
            continue
        path = target[1:] if target.startswith("/") else posixpath.join("xl", target)
        ranges: list[Bounds] = []
        with archive.open(posixpath.normpath(path)) as source:
            for _, element in ElementTree.iterparse(source):
                if element.tag == f"{MAIN_NS}mergeCell":
                    bounds = range_boundaries(element.get("ref", ""))
                    if None not in bounds:
                        ranges.append(bounds)  # type: ignore[arg-type]
                element.clear()
        result[sheet.get("name", "")] = tuple(ranges)
    return result


def apply_merges(sheet: Sheet, ranges: tuple[Bounds, ...]) -> Sheet:
    """Copy the value of each vertical merge down its first column, on rows that have content.

    Excel keeps a merged value in the top-left cell only; a company merged over three people's
    rows belongs to each of them. Horizontal merges copy nothing (the value stays in the first
    column). Copied cells are flagged for a row diagnostic.
    """
    if not ranges:
        return sheet
    rows = {row.number: row for row in sheet.rows}
    for min_col, min_row, _, max_row in ranges:
        top = rows.get(min_row)
        value = top.value(min_col - 1) if top else None
        if is_blank(value):
            continue
        for number in range(min_row + 1, max_row + 1):
            row = rows.get(number)
            if row is None or not is_blank(row.value(min_col - 1)):
                continue
            values = list(row.values) + [None] * (min_col - len(row.values))
            values[min_col - 1] = value
            rows[number] = replace(row, values=tuple(values), copied=row.copied | {min_col - 1})
    ordered = tuple(rows[row.number] for row in sheet.rows)
    width = max([sheet.width, *(len(row.values) for row in ordered)])
    return replace(sheet, rows=ordered, width=width, merged_ranges=len(ranges))


# --- CSV ----------------------------------------------------------------------------------------


def decode_csv(content: bytes) -> tuple[str, str]:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return content.decode("utf-16"), "utf-16"
        except UnicodeDecodeError:
            raise ImportRejectedError(DiagnosticCode.FILE_ENCODING_UNKNOWN) from None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding), encoding.removesuffix("-sig")
        except UnicodeDecodeError:
            continue
    raise ImportRejectedError(DiagnosticCode.FILE_ENCODING_UNKNOWN)


def sniff_delimiter(text: str) -> str:
    """The delimiter giving the most records of one consistent width (> 1) in the first ones."""

    def score(delimiter: str) -> tuple[int, int]:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        try:
            widths = [len(record) for record in islice(reader, SNIFF_RECORDS) if any(record)]
        except csv.Error:
            return (0, 0)
        if not widths:
            return (0, 0)
        width, count = Counter(widths).most_common(1)[0]
        return (count, width) if width > 1 else (0, 0)

    return max(CSV_DELIMITERS, key=score)


def read_csv(name: str, text: str, delimiter: str, limits: ImportLimits) -> Sheet:
    rows: list[SheetRow] = []
    try:
        records = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        for number, record in enumerate(records, start=1):
            cells = trimmed(tuple(value or None for value in record))
            if cells:
                check_bounds(name, len(rows) + 1, len(cells), limits)
                rows.append(SheetRow(number, cells))
    except csv.Error:
        raise ImportRejectedError(DiagnosticCode.FILE_CORRUPT) from None
    return Sheet(name, tuple(rows), max((len(row.values) for row in rows), default=0))
