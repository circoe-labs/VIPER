"""XLSX writer of the Excel export (Task 10): readable, typed, formula-free and deterministic.

- **Readable**: bold header row, frozen under the header, autofilter over the data, column widths
  from the specification.
- **Typed**: dates are real Excel dates (Europe/Paris calendar), integers are numbers, texts are
  text cells; `CODE` columns (phones, SIREN, SIRET, postal codes, ids) use the Text format so a
  leading zero or `+` survives an edit in Excel.
- **Formula-free**: every text is written as a string cell — openpyxl would otherwise store a text
  starting with `=` as a formula and `#N/A` as an error. A text Excel would read as a formula when
  re-typed (the shared rule of `app.core.spreadsheet`, as for the explorer's CSV) also gets the
  `quotePrefix` style, so it stays text after an edit; its value is kept exactly.
- **Deterministic**: no clock, no randomness — the creation date is an argument and every zip entry
  carries it, so the same data and date always give the same bytes.

Characters XML cannot hold (control characters other than tab and line breaks) are dropped and a
text longer than Excel's 32 767-character cell limit is cut, as Excel itself would refuse them.
"""

from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, time
from io import BytesIO
from typing import IO, Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from openpyxl.writer.excel import ExcelWriter

from app.core.business_time import BUSINESS_TIMEZONE
from app.core.spreadsheet import looks_like_formula
from app.services.exports.spec import Column, Kind, Value

EXCEL_MAX_TEXT = 32_767
DATE_FORMAT = "dd/mm/yyyy"
DATETIME_FORMAT = "dd/mm/yyyy hh:mm"
TEXT_FORMAT = "@"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F2A37")
HEADER_BORDER = Border(bottom=Side(style="thin", color="14B8A6"))
HEADER_ALIGNMENT = Alignment(vertical="center")


@dataclass(frozen=True, slots=True)
class SheetContent:
    key: str
    name: str
    columns: Sequence[Column[Any]]
    rows: Iterable[Sequence[Value]]


def write_workbook(
    sheets: Iterable[SheetContent], *, created: datetime
) -> tuple[bytes, dict[str, int]]:
    """The XLSX bytes and the number of data rows written per sheet (by key)."""
    book = Workbook(write_only=True)
    moment = created.astimezone(UTC).replace(tzinfo=None, microsecond=0)
    book.properties.creator = "VIPER"
    book.properties.title = "Export VIPER"
    book.properties.created = moment
    book.properties.modified = moment
    counts = {sheet.key: _write_sheet(book, sheet) for sheet in sheets}
    buffer = BytesIO()
    with _StableZip(buffer, "w", ZIP_DEFLATED, stamp=moment) as archive:
        ExcelWriter(book, archive).save()
    return buffer.getvalue(), counts


def _write_sheet(book: Workbook, sheet: SheetContent) -> int:
    worksheet = book.create_sheet(sheet.name)
    assert isinstance(worksheet, WriteOnlyWorksheet)
    # Widths and panes must be set before the first row of a write-only sheet. (The stubs of
    # `WriteOnlyWorksheet` miss `column_dimensions` and `auto_filter`, which it has at runtime.)
    for index, column in enumerate(sheet.columns, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = column.width  # type: ignore[attr-defined]
    worksheet.freeze_panes = "A2"
    worksheet.append([_header(worksheet, column.header) for column in sheet.columns])
    count = 0
    for values in sheet.rows:
        cells = zip(sheet.columns, values, strict=True)
        worksheet.append([_cell(worksheet, column.kind, value) for column, value in cells])
        count += 1
    last = get_column_letter(len(sheet.columns))
    worksheet.auto_filter.ref = f"A1:{last}{count + 1}"  # type: ignore[attr-defined]
    return count


def _header(worksheet: WriteOnlyWorksheet, text: str) -> Cell:
    cell = _text_cell(worksheet, text)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = HEADER_BORDER
    cell.alignment = HEADER_ALIGNMENT
    return cell


def _cell(worksheet: WriteOnlyWorksheet, kind: Kind, value: Value) -> Cell | None:
    if value is None or value == "":
        return None
    match kind:
        case Kind.DATE:
            assert isinstance(value, date)
            day = _local(value).date() if isinstance(value, datetime) else value
            return _dated(worksheet, datetime.combine(day, time()))
        case Kind.DATETIME:
            assert isinstance(value, datetime)
            moment = _local(value)
            return _dated(
                worksheet, moment, DATE_FORMAT if moment.time() == time() else DATETIME_FORMAT
            )
        case Kind.CODE:
            cell = _text_cell(worksheet, str(value))
            cell.number_format = TEXT_FORMAT
            return cell
        case _ if isinstance(value, str):
            return _text_cell(worksheet, value)
        case _:
            assert isinstance(value, int | float)  # booleans included
            return WriteOnlyCell(worksheet, value)


def _local(moment: datetime) -> datetime:
    """Excel has no time zones: the Europe/Paris wall-clock time, naive."""
    return moment.astimezone(BUSINESS_TIMEZONE).replace(tzinfo=None) if moment.tzinfo else moment


def _dated(
    worksheet: WriteOnlyWorksheet, value: datetime, number_format: str = DATE_FORMAT
) -> Cell:
    cell = WriteOnlyCell(worksheet, value)
    cell.number_format = number_format
    return cell


def _text_cell(worksheet: WriteOnlyWorksheet, value: str) -> Cell:
    text = ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_MAX_TEXT]
    cell = WriteOnlyCell(worksheet, text)
    cell.data_type = "s"  # never a formula (`=…`) nor an error value (`#N/A`)
    if looks_like_formula(text):
        cell.quotePrefix = True
    return cell


class _StableZip(ZipFile):
    """Every entry stamped with the export's date (instead of the current time or a temporary
    file's modification time), so identical content always gives identical bytes."""

    def __init__(self, file: IO[bytes], mode: str, compression: int, *, stamp: datetime) -> None:
        super().__init__(file, mode, compression)  # type: ignore[call-overload]
        self.stamp = stamp.timetuple()[:6]

    def writestr(
        self,
        zinfo_or_arcname: str | ZipInfo,
        data: Any,
        compress_type: int | None = None,
        compresslevel: int | None = None,
    ) -> None:
        if isinstance(zinfo_or_arcname, str):
            zinfo_or_arcname = ZipInfo(zinfo_or_arcname, date_time=self.stamp)
            zinfo_or_arcname.compress_type = self.compression
        super().writestr(zinfo_or_arcname, data, compress_type, compresslevel)

    def write(
        self,
        filename: Any,
        arcname: Any = None,
        compress_type: int | None = None,
        compresslevel: int | None = None,
    ) -> None:
        with open(filename, "rb") as source:
            self.writestr(str(arcname or filename), source.read(), compress_type, compresslevel)
